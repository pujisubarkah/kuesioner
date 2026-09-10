import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_paper():
    doc = docx.Document()

    # Page Margins (Standard 1 inch / 72 pt)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # TITLE
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(14)
    title_p.paragraph_format.line_spacing = 1.15
    title_run = title_p.add_run("Context-Aware Behavioral Analytics for Adaptive Public Sector Training: An Empirical Baseline and Framework Design")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64) # Navy

    # ABSTRACT BOX
    abs_heading = doc.add_paragraph()
    abs_heading.paragraph_format.space_before = Pt(6)
    abs_heading.paragraph_format.space_after = Pt(4)
    r = abs_heading.add_run("Abstract")
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    abs_p = doc.add_paragraph()
    abs_p.paragraph_format.space_before = Pt(0)
    abs_p.paragraph_format.space_after = Pt(8)
    abs_p.paragraph_format.line_spacing = 1.15
    abs_run = abs_p.add_run(
        "Digital transformation in public governance has mandated continuous competency development for civil servants (PNS/ASN), predominantly delivered via asynchronous Learning Management Systems (LMS). However, current platforms predominantly suffer from context blindness—failing to account for situational realities such as fragmented working hours, shifting administrative priorities, multi-device usage, infrastructural disparities in remote/3T regions, and compliance-driven ghost learning. This paper presents an empirically grounded framework for Context-Aware Behavioral Analytics (CABA) specifically designed for public sector professional training. Based on a formative empirical survey of Indonesian civil servants across diverse administrative tiers and geographical regions, we establish a baseline of real-world learning constraints, including prevalent operational multitasking, connectivity bottlenecks, and modality preferences. Building upon these empirical insights, we formulate a four-layer CABA architecture that bridges multi-dimensional context sensing (temporal workload, environmental infrastructure, organizational tier, and non-invasive behavioral telemetry) with an adaptive pedagogical decision engine. The framework dynamically delivers supportive, non-punitive interventions, such as workload-sensitive micro-chunking, active engagement verification stop-gates, and bandwidth-resilient modality fallbacks. This study provides both foundational empirical evidence and an architectural blueprint for next-generation government e-learning ecosystems that optimize authentic competency acquisition while respecting public servants' operational workload."
    )
    abs_run.font.size = Pt(10)
    abs_run.font.italic = True

    kw_p = doc.add_paragraph()
    kw_p.paragraph_format.space_before = Pt(0)
    kw_p.paragraph_format.space_after = Pt(18)
    r1 = kw_p.add_run("Keywords: ")
    r1.font.bold = True
    r1.font.size = Pt(10)
    r2 = kw_p.add_run("Learning Analytics, Context-Aware Computing, Public Sector Training, Adaptive Learning, Civil Service Competency, Empirical Baseline, Ghost Learning, Andragogy.")
    r2.font.italic = True
    r2.font.size = Pt(10)

    # Helper function for headings
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
        return h

    def add_p(text, bold_prefix=None, space_after=6, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
        r = p.add_run(text)
        r.font.italic = italic
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
        p.add_run(text)
        return p

    # SECTION I
    add_h1("I. Introduction")
    add_h2("1.1 Background and Motivation")
    add_p("In modern public administration, continuous competency development for civil servants (Pegawai Negeri Sipil/PNS or Aparatur Sipil Negara/ASN) is both a governance imperative and a regulatory mandate, often operationalized through compulsory annual training hours and competency-based certification requirements. Governments worldwide have accelerated the adoption of asynchronous Massive Open Online Courses (MOOCs), digital learning portals, and Learning Management Systems (LMS) to upskill large numbers of public sector employees efficiently. While digital delivery helps overcome geographic, logistical, and cost barriers, the pedagogical efficacy of these platforms remains heavily contested, especially when learning is treated primarily as a compliance requirement rather than as a process of authentic competency acquisition (Gašević et al., 2015; Siemens and Baker, 2012).")
    add_p("Unlike conventional higher education students who often operate within structured academic timetables, civil servants engage in online learning under strict adult learning and andragogical constraints. They must balance continuous professional education with daily public service delivery, administrative deadlines, shifting task priorities, and unexpected state duties. Consequently, their learning behaviors are deeply situated within workplace routines, organizational expectations, and operational pressures (Brown et al., 1989; Knowles, 1984).")

    add_h2("1.2 Problem Statement and the Context Blindness Gap")
    add_p("Current public sector e-learning ecosystems often suffer from context blindness. Standard systems treat online learning interactions in isolation, failing to account for three critical operational constraints:")
    add_bullet(" Standard modules frequently consist of lengthy video lectures or monolithic reading materials that clash with the fragmented daily schedules of public servants, leading to cognitive overload, rushed completion, or high dropout rates.", "1. Mismatch between Time Fragmentation and Monolithic Content:")
    add_bullet(" Mandatory training quotas frequently incentivize superficial compliance or ghost learning behaviors—such as playing training videos in background tabs while attending to official duties, skipping reading materials, or rapidly guessing post-tests without meaningful cognitive processing.", "2. Compliance-Driven Ghost Learning Behaviors:")
    add_bullet(" Public servants are stationed across widely disparate technological landscapes—from central ministry headquarters with gigabit fiber to Disadvantaged, Frontier, and Outermost regions (Wilayah 3T) with unstable cellular networks and multi-device dependencies.", "3. Infrastructural and Spatial Heterogeneity:")
    add_p("Traditional Learning Analytics (LA) relies primarily on static behavioral logs, such as total watch time, login frequency, completion status, and post-test scores. However, analyzing learner behavior in isolation from situational context produces biased interpretations. For example, a civil servant watching a video at 1.5× speed at 10:00 AM on an office desktop may indicate high workload and time pressure during official service hours, whereas the same behavior at 9:00 PM on a mobile phone may reflect self-paced review. The problem is not the absence of behavioral telemetry, but the lack of situational awareness required to interpret such data pedagogically.")

    add_h2("1.3 Research Contributions")
    add_p("To bridge this gap, this paper introduces an empirically grounded Context-Aware Behavioral Analytics (CABA) framework tailored for public sector professional training. Its primary contributions are threefold:")
    add_bullet(" Formulating a multi-regional baseline of Indonesian civil servants' learning realities, detailing empirical workplace interruptions, device ecologies, connectivity constraints, and behavioral multitasking patterns across central and regional government bodies.", "1. Empirical Baseline of Public Sector Learning Realities:")
    add_bullet(" Defining multi-modal context indicators—including temporal workload, environmental and infrastructural constraints, organizational roles, and non-invasive telemetry proxies.", "2. Taxonomy of Contextual and Behavioral Dimensions:")
    add_bullet(" Developing an end-to-end conceptual pipeline connecting passive context sensing, probabilistic state inference (mitigating ghost learning and cognitive fatigue), an andragogy-centered pedagogical decision engine, and supportive, non-punitive intervention delivery (e.g., dynamic micro-chunking and adaptive modality fallback).", "3. Four-Layer Conceptual Architecture & Intervention Strategies:")

    add_h2("1.4 Paper Organization")
    add_p("The remainder of this paper is organized as follows. Section II discusses the theoretical foundations and related work. Section III presents the empirical baseline derived from our formative civil service survey. Section IV presents the proposed CABA framework architecture and inference logic. Section V provides operational scenarios illustrating how CABA functions across diverse workplace contexts. Section VI discusses ethical considerations, non-punitive governance, and system limitations. Finally, Section VII concludes the paper and outlines a future roadmap.")

    # SECTION II
    add_h1("II. Theoretical Foundations and Related Work")
    add_h2("2.1 Theoretical Grounding: Situated Cognition, Self-Regulated Learning, and Andragogy")
    add_p("The design of the Context-Aware Behavioral Analytics (CABA) framework is grounded in three complementary pedagogical and cognitive paradigms: situated cognition, self-regulated learning, and andragogy. Situated Cognition Theory argues that knowledge is not detached from the conditions in which it is acquired and used, but is shaped by activity, context, and culture (Brown et al., 1989). In public sector professional training, this perspective is particularly relevant because civil servants acquire competencies while embedded in authentic administrative routines, organizational hierarchies, service obligations, and task disruptions.")
    add_p("Self-Regulated Learning (SRL) provides a second foundation. Online professional development requires learners to manage goals, time, learning environments, help-seeking behaviors, task strategies, and self-evaluation. Prior work shows that learners differ substantially in their capacity to regulate learning, and that disorganized self-regulation is associated with poorer achievement (Barnard-Brak et al., 2010). In mandatory government training, where compliance targets may overshadow intrinsic motivation, adaptive systems need mechanisms to detect weak regulation patterns and provide timely scaffolding.")
    add_p("Andragogy further strengthens the public sector orientation of CABA. Adult learners are typically self-directed, experience-rich, problem-oriented, and motivated by the relevance of learning to their work roles (Knowles, 1984). Effective online training should connect learning materials with real administrative problems, acknowledge prior professional experience, and provide flexible pathways that respect workload constraints.")

    add_h2("2.2 Behavioral Analytics and Engagement Estimation")
    add_p("The evolution of Learning Analytics (LA) and Educational Data Mining (EDM) has shifted from post-hoc reporting toward increasingly granular behavioral tracking (Siemens and Baker, 2012). Contemporary platforms capture clickstream events, resource navigation sequences, dwell time, video playback behavior, quiz attempts, and response latency. For public sector training, engagement estimation must prioritize non-invasive behavioral proxies rather than hardware-heavy or privacy-invasive sensing methods (such as eye-tracking or webcam monitoring). Indicators such as tab-focus changes, repeated video scrubbing, abnormal playback acceleration, and inconsistent resource sequencing help identify disengagement without compromising learner privacy.")

    add_h2("2.3 Contextual Constraints and Infrastructural Heterogeneity in the Public Sector")
    add_p("A critical gap in existing adaptive learning literature is the assumption that learners operate under relatively stable digital infrastructure. In public sector ecosystems, learners are geographically dispersed. In Indonesia, civil servants are located across central ministries, provincial agencies, municipal offices, and remote 3T regions. These settings vary significantly in bandwidth stability, device quality, and opportunities for uninterrupted study. Context awareness must therefore incorporate infrastructural context alongside workplace temporal context.")

    add_h2("2.4 Synthesis and Comparative Analysis")
    add_p("Table 1 summarizes the comparison between traditional LMS learning analytics, general academic adaptive systems, and the proposed CABA framework.")

    # TABLE 1
    t1 = doc.add_table(rows=5, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    
    headers_t1 = ["Dimension", "Traditional LMS Analytics", "General Adaptive E-Learning", "Proposed CABA Framework"]
    for i, title in enumerate(headers_t1):
        cell = t1.cell(0, i)
        set_cell_background(cell, "1F3864")
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_t1 = [
        ("Primary Target Audience", "K-12 and higher education students", "General academic or lifelong learners", "Civil servants (PNS/ASN) and public sector professionals"),
        ("Context Dimensions", "Minimal; primarily log timestamps", "Academic schedules and prior test scores", "Multidimensional: temporal workload blocks, device/bandwidth constraints, 3T regional infrastructure, institutional tier"),
        ("Behavioral Telemetry", "Page visits, login duration, test completion", "Navigation sequences and quiz response times", "In-video interactions, tab-focus dynamics, response latency, idle periods, and anti-ghost-learning telemetry"),
        ("Adaptive Interventions", "Static reports for instructors", "Automated difficulty adjustment and content sequencing", "Situated interventions: office-hour micro-chunking, engagement checkpoints, bandwidth-resilient fallback, role scaffolding")
    ]

    for row_idx, row_data in enumerate(data_t1, start=1):
        for col_idx, text in enumerate(row_data):
            cell = t1.cell(row_idx, col_idx)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F9")
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True

    caption1 = doc.add_paragraph()
    caption1.paragraph_format.space_before = Pt(4)
    caption1.paragraph_format.space_after = Pt(14)
    r = caption1.add_run("Table 1. Comparison between Traditional Learning Analytics, General Adaptive Systems, and CABA.")
    r.font.size = Pt(9)
    r.font.italic = True

    # SECTION III (NEW EMPIRICAL BASELINE)
    add_h1("III. Empirical Baseline: Formative Survey on Public Sector Learning Realities")
    add_h2("3.1 Survey Instrument, Sampling, and Demographic Administration")
    add_p("To ground the architectural design of CABA in empirical reality, a comprehensive formative questionnaire was administered to Indonesian civil servants across diverse administrative jurisdictions. The instrument comprised nine operational modules (Sections A through I):")
    add_bullet(" Capturing geographical tiers (provinces, regencies/cities, rural, border, island, and 3T zones), institutional levels (Central Ministries/Agencies, Provincial Governments, Municipal/Regency Governments), age brackets, and years of civil service tenure.", "Section A (Participant Demographics & Regional Profiles):")
    add_bullet(" Assessing historical participation in government MOOCs, Latsar (pre-service training), leadership academies, functional training, and digital platforms (e.g., ASN Unggul, Kolabjar, LMS Instansi, Moodle).", "Section B (E-Learning Participation History):")
    add_bullet(" Measuring primary devices (desktop, laptop, smartphone, tablet, multi-device switching), physical learning locations, connectivity sources, frequency of technical disruptions, and open qualitative experiences.", "Section C (Hardware Devices & Connectivity Ecology):")
    add_bullet(" Measuring concurrent public service duties during training, frequency of duty-mandated task abandonment, time availability, and physical workplace mobility (Likert 1–5).", "Section D (Workplace Environment & Operational Workload):")
    add_bullet(" Evaluating content accessibility, video file size barriers, job relevance, and modality preferences (Text, Video, Audio, Infographics, Simulations, Case Studies).", "Section E (Learning Content & Modality Preferences):")
    add_bullet(" Evaluating attention focus, peer/facilitator interactivity, and primary causes of disengagement (workload, fatigue, connectivity, platform friction).", "Section F (Interaction Dynamics & Disengagement Triggers):")
    add_bullet(" In-depth qualitative reflections uncovering systemic bottlenecks, regional disparity conditions, and institutional blind spots.", "Sections G & H (Open Reflection & Policy Blind Spots):")
    add_bullet(" Matrix evaluation of virtual classroom behaviors, including webcam usage, background multitasking, mobile attendance, and connection stability.", "Section I (Synchronous Video Session Realities):")

    add_h2("3.2 Key Empirical Findings and Ground Truth")
    add_p("The empirical baseline revealed four defining operational characteristics of public sector online learning:")
    add_p("A substantial majority of civil servants report that online learning is conducted concurrently with active public service obligations. Learners are frequently forced to abandon learning sessions midway due to urgent citizen service requests or superior directives. Time availability is fragmented into short, unpredictable intervals, directly challenging monolithic module designs.", bold_prefix="1. Pervasive Task Fragmentation during Office Hours: ")
    add_p("Learners stationed outside major metropolitan hubs and in 3T/border areas rely heavily on mobile data connections and multi-device switching. High-bitrate video streams produce severe buffering, high data quota consumption, and session timeouts, causing severe accessibility disparities.", bold_prefix="2. Severe Infrastructural & Bandwidth Constraints: ")
    add_p("When training is treated as a rigid compliance metric without workload accommodation, learners adopt rational coping mechanisms: running videos in background tabs while typing office documents, attending synchronous webinars with cameras and audio disabled while attending field assignments, and skipping straight to evaluations.", bold_prefix="3. Rationalized Ghost Learning as a Coping Mechanism: ")
    add_p("Learners overwhelmingly express higher comprehension and task-resilience when provided with concise, multi-modal content formats (infographics, structured executive summaries, and downloadable audio/case studies) rather than 45-minute continuous video lectures.", bold_prefix="4. Demand for Modality Flexibility and Micro-Learning: ")

    add_h2("3.3 Mapping Empirical Findings to CABA Framework Components")
    add_p("Table 2 outlines how specific survey findings directly justify the four operational layers of the CABA architecture.")

    # TABLE 2 (SURVEY MAPPING)
    t2 = doc.add_table(rows=5, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    headers_t2 = ["Empirical Survey Dimension", "Observed Field Phenomenon", "Addressed CABA Layer", "Target System Adaptation"]
    for i, title in enumerate(headers_t2):
        cell = t2.cell(0, i)
        set_cell_background(cell, "1F3864")
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_t2 = [
        ("Section D (Workload) & Section I (Sync)", "Severe office duty multitasking; off-camera webinar attendance; abrupt session exits", "Layer 1 (Temporal) & Layer 2 (State S2)", "Workload-aware time tracking and auto-saving session bookmarks upon tab switching"),
        ("Section C (Connectivity) & Section A (3T)", "Packet loss, video buffering, and quota exhaustion in remote regencies", "Layer 1 (Infra) & Layer 2 (State S4)", "Automatic network sensing and seamless fallback to text-first / audio-first modules"),
        ("Section F (Focus) & Section I (Sync)", "Background tab playback; rapid guessing without reading material", "Layer 2 (State S1) & Layer 3 (Decision)", "Active verification stop-gates and interactive 1-minute scenario checkpoints"),
        ("Section E (Modality Preferences)", "Strong preference for micro-content, infographics, and situational case studies", "Layer 4 (Delivery Interface)", "Dynamic micro-chunking into 3–5 minute digestible learning units")
    ]

    for row_idx, row_data in enumerate(data_t2, start=1):
        for col_idx, text in enumerate(row_data):
            cell = t2.cell(row_idx, col_idx)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F9")
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True

    caption2 = doc.add_paragraph()
    caption2.paragraph_format.space_before = Pt(4)
    caption2.paragraph_format.space_after = Pt(14)
    r = caption2.add_run("Table 2. Empirical Grounding Matrix: Survey Dimensions Mapped to CABA Architectural Layers.")
    r.font.size = Pt(9)
    r.font.italic = True

    # SECTION IV (CABA ARCHITECTURE)
    add_h1("IV. The Proposed CABA Framework Architecture")
    add_h2("4.1 High-Level Architecture Overview")
    add_p("The Context-Aware Behavioral Analytics (CABA) framework is structured as a four-tier pipeline designed to continuously capture, infer, decide, and deliver situated interventions for public sector learners. The framework connects non-invasive telemetry collection with contextual inference and pedagogical decision-making so that adaptive interventions are triggered not by behavior alone, but by behavior interpreted within workplace, infrastructural, and organizational conditions.")

    # Embed Figure 1
    if os.path.exists("data/figures/figure1.png"):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig1.paragraph_format.space_before = Pt(10)
        p_fig1.paragraph_format.space_after = Pt(4)
        run_fig1 = p_fig1.add_run()
        run_fig1.add_picture("data/figures/figure1.png", width=Inches(5.8))
        
        caption_fig1 = doc.add_paragraph()
        caption_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_fig1.paragraph_format.space_before = Pt(2)
        caption_fig1.paragraph_format.space_after = Pt(12)
        r = caption_fig1.add_run("Figure 1. The Four-Layer Architecture of the Context-Aware Behavioral Analytics (CABA) Framework.")
        r.font.size = Pt(9)
        r.font.italic = True

    add_p("The four layers are: Layer 1, Context Sensing and Behavioral Telemetry; Layer 2, Context Fusion and Learner State Inference; Layer 3, Adaptive Pedagogical Decision Engine; and Layer 4, Situated Intervention and Delivery Interface. These layers operate iteratively, allowing the LMS to update learner-state estimates and adapt interventions during an ongoing learning session.")

    add_h2("4.2 Layer 1: Context Sensing and Behavioral Telemetry")
    add_p("Layer 1 non-invasively captures real-time data streams without requiring biometric sensors, thereby supporting institutional privacy requirements and reducing implementation barriers in government LMS environments. At time t, the real-time context vector is denoted as:")

    p_eq = doc.add_paragraph()
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(6)
    p_eq.paragraph_format.space_after = Pt(6)
    r_eq = p_eq.add_run("C(t) = ⟨ Twork(t), Einfra(t), Ogov, Btele(t) ⟩")
    r_eq.font.bold = True
    r_eq.font.size = Pt(11)

    add_p("where Twork(t) represents temporal and workload context, Einfra(t) represents infrastructural and environmental context, Ogov represents organizational and administrative context, and Btele(t) represents behavioral telemetry.")
    add_bullet(" thour captures time-of-day categories, such as active public service hours, overtime/evening learning, and late-night learning; Δtsession captures continuous session duration and inter-session gaps.", "Temporal and workload context: ")
    add_bullet(" βnet estimates network bandwidth and latency, including potential 3T or remote-office constraints; δdev identifies device category, such as desktop workstation, mobile phone, or tablet.", "Infrastructural and environmental context: ")
    add_bullet(" Rtier represents civil service rank or role tier, such as administrative, specialist, or executive; Qjp represents the remaining mandatory training quota.", "Organizational and administrative context: ")
    add_bullet(" τfocus captures LMS window or tab focus ratio; νplay captures video playback speed; Nscrub captures non-sequential timeline scrubbing; and Tdwell compares dwell time per learning unit with an expected benchmark.", "Behavioral telemetry: ")

    add_h2("4.3 Layer 2: Context Fusion and Learner State Inference")
    add_p("Layer 2 aggregates raw telemetry features into discrete and interpretable learner cognitive and engagement states. Rather than treating any single behavior as conclusive, the framework combines threshold heuristics, temporal sequences, and contextual rules to infer learner states probabilistically. Four primary states are formulated:")
    add_bullet(" Inferred when τfocus is below a predefined focus threshold, or when high playback speed, excessive scrubbing, and unusually short dwell time occur together. It indicates that the learner may be running training materials in the background primarily to fulfill mandatory hours.", "1. S1: Ghost Learning or Passive Compliance.")
    add_bullet(" Inferred when learning occurs during active public service hours, usually on an office desktop, with frequent session interruptions. It indicates an authentic learning attempt constrained by administrative duties.", "2. S2: Fragmented Office-Hour Learning.")
    add_bullet(" Inferred when repeated video rewinds, dwell time exceeding twice the expected benchmark, and low quiz performance co-occur. It suggests that the learner may be struggling with complex technical regulations or procedural materials.", "3. S3: Cognitive Overload or Concept Struggling.")
    add_bullet(" Inferred when network bandwidth falls below a threshold or when video buffering stall ratio is high. It indicates connectivity limitations common in regional or 3T learning environments.", "4. S4: Bandwidth-Constrained Access.")

    add_h2("4.4 Layer 3: Adaptive Pedagogical Decision Engine")
    add_p("The Decision Engine applies an andragogy-centered mapping function F: S × Ogov → A to select an appropriate pedagogical intervention A. The selected intervention is represented as:")

    p_eq2 = doc.add_paragraph()
    p_eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq2.paragraph_format.space_before = Pt(6)
    p_eq2.paragraph_format.space_after = Pt(6)
    r_eq2 = p_eq2.add_run("A* = arg max_a∈A  U(a | Sk, Ogov)")
    r_eq2.font.bold = True
    r_eq2.font.size = Pt(11)

    add_p("where U represents a pedagogical utility function that aims to maximize competency retention while minimizing workplace disruption. Table 3 presents the decision mapping matrix.")

    # TABLE 3
    t3 = doc.add_table(rows=5, cols=4)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)

    headers_t3 = ["Inferred State", "Contextual Triggers", "Target Adaptation Policy", "Pedagogical Rationale"]
    for i, title in enumerate(headers_t3):
        cell = t3.cell(0, i)
        set_cell_background(cell, "1F3864")
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_t3 = [
        ("S1: Ghost Learning", "Low tab-focus ratio; video running in background; high acceleration", "Active verification stop-gate: pause playback and inject a 1-question scenario checkpoint", "Converts passive compliance into active cognitive processing"),
        ("S2: Fragmented Office-Hour Learning", "Office hours; desktop access; frequent interruptions", "Dynamic micro-chunking: break long modules into short digestible units", "Accommodates adult time constraints without abandoning competency goals"),
        ("S3: Cognitive Struggling", "Repeated rewinds; extended dwell time; low quiz scores", "Scaffolded regulatory clarification: provide infographics, simplified examples, hints", "Reduces cognitive overload and supports situated understanding"),
        ("S4: Low Bandwidth / 3T Access", "Low bandwidth; high packet loss; repeated buffering", "Adaptive modality fallback: switch to low-bitrate slides, audio, or text-first mode", "Improves accessibility and learning continuity across unequal infrastructures")
    ]

    for row_idx, row_data in enumerate(data_t3, start=1):
        for col_idx, text in enumerate(row_data):
            cell = t3.cell(row_idx, col_idx)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F9")
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True

    caption3 = doc.add_paragraph()
    caption3.paragraph_format.space_before = Pt(4)
    caption3.paragraph_format.space_after = Pt(14)
    r = caption3.add_run("Table 3. Adaptive Decision Mapping Matrix.")
    r.font.size = Pt(9)
    r.font.italic = True

    add_h2("4.5 Layer 4: Situated Intervention and Delivery Interface")
    add_p("Layer 4 renders the selected intervention directly into the learner interface. Interventions are strictly designed to be supportive rather than punitive. For example, instead of penalizing tab-switching, the system preserves the learner's timestamp and provides a brief recap upon return. When infrastructural constraints are detected, the platform automatically presents bandwidth-resilient formats. Role differentiation tailors case studies to structural leaders (policy formulation) or technical staff (procedural workflows).")

    # SECTION V (OPERATIONAL SCENARIOS)
    add_h1("V. Operational Scenarios and Illustrative Walkthrough")
    add_p("To demonstrate the conceptual validity and real-world applicability of CABA, this section presents two illustrative operational scenarios grounded in Indonesian civil service realities.")

    # Embed Figure 2
    if os.path.exists("data/figures/figure2.png"):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig2.paragraph_format.space_before = Pt(10)
        p_fig2.paragraph_format.space_after = Pt(4)
        run_fig2 = p_fig2.add_run()
        run_fig2.add_picture("data/figures/figure2.png", width=Inches(5.8))
        
        caption_fig2 = doc.add_paragraph()
        caption_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_fig2.paragraph_format.space_before = Pt(2)
        caption_fig2.paragraph_format.space_after = Pt(12)
        r = caption_fig2.add_run("Figure 2. Illustrative Workflow of CABA Inference and Intervention across Operational Learning Scenarios.")
        r.font.size = Pt(9)
        r.font.italic = True

    add_h2("5.1 Scenario 1: Mitigating Ghost Learning During Peak Office Hours")
    add_p("A civil servant in an IT support role accesses a mandatory 45-minute training module on public procurement at 10:15 AM using a desktop workstation in a central ministry. The telemetry stream indicates a tab-focus ratio of τfocus = 0.25 (active for only a quarter of the interval) while playback speed is set to νplay = 1.75×. The context fusion layer classifies this pattern as S1: Ghost Learning superimposed on S2: Fragmented Office-Hour Learning. The decision engine avoids punitive measures; instead, playback is paused gracefully, the timestamp is bookmarked, and the remaining content is transformed into a 3-minute interactive micro-case. When returning to the LMS tab, the learner receives a short contextual recap and a practical decision check.")

    add_h2("5.2 Scenario 2: Adaptive Scaffolding in a Bandwidth-Constrained 3T Outpost")
    add_p("A civil servant stationed in a remote frontier district accesses an ethics module via a smartphone over an unstable cellular network. The telemetry stream registers throughput below βnet < 120 kbps and a video stall ratio exceeding 0.45. The context fusion layer infers S4: Bandwidth-Constrained Access. The system immediately activates adaptive modality fallback, replacing heavy video streams with structured low-bandwidth text summaries, visual infographics, and optional compressed audio snippets. Quiz items are pre-cached locally to prevent submission timeout errors.")

    # SECTION VI (DISCUSSION)
    add_h1("VI. Discussion, Ethical Considerations, and Limitations")
    add_h2("6.1 Non-Punitive Analytics versus Workplace Surveillance")
    add_p("In public sector environments, behavioral analytics can easily be perceived as intrusive workplace surveillance if governance boundaries are unclear. CABA is explicitly positioned as a pedagogical enablement system rather than a disciplinary monitoring tool. Its purpose is to support learning continuity and reduce cognitive burden, not to evaluate employee productivity. Two principles guide this: (1) Data Minimization—capturing only coarse-grained telemetry without logging personal browsing, keystrokes, or webcams; and (2) Transparent Learner Agency—informing learners why an adaptation occurred through clear UI explanations.")

    add_h2("6.2 Mitigating Over-Intervention and Alert Fatigue")
    add_p("To avoid distracting learners, CABA incorporates an intervention cooldown mechanism, denoted as Δtcooldown ≥ 15 minutes. Nudges and stop-gates are deployed sparingly and only when disengagement or bottlenecks are detected with high statistical confidence.")

    add_h2("6.3 Limitations and Threats to Validity")
    add_bullet(" Decision thresholds (θfocus, θbandwidth, θscrub) require empirical parameter tuning against real LMS transaction logs.", "Threshold Heuristics:")
    add_bullet(" While the empirical baseline provides strong exploratory validity, full behavioral validation requires long-term quasi-experimental pilot deployments across diverse government cohorts.", "Field Trial Scope:")
    add_bullet(" Operational deployment requires formal data consent protocols and institutional algorithmic transparency guidelines.", "Data Governance:")

    # SECTION VII (CONCLUSION)
    add_h1("VII. Conclusion and Future Roadmap")
    add_p("This paper presented an empirically grounded Context-Aware Behavioral Analytics (CABA) framework for public sector e-learning. By synthesizing adult learning theory, situated cognition, and real-world survey evidence from Indonesian civil servants, CABA addresses the critical context blindness of conventional LMS platforms. The four-tier architecture connects multi-dimensional context sensing with an andragogical decision engine to deliver non-punitive adaptations, including active verification stop-gates, dynamic micro-chunking, and bandwidth-resilient fallbacks.")
    add_p("Future work will proceed in two directions: First, developing a reference plugin implementation for widely used government LMS platforms (such as Moodle). Second, executing a quasi-experimental pilot study across civil service training cohorts to evaluate the framework's impact on knowledge retention and engagement quality.")

    # REFERENCES
    add_h1("References")
    refs = [
        "Barnard-Brak, L., Lan, W. Y., and Paton, V. O. (2010). Profiles in self-regulated learning in the online learning environment. The International Review of Research in Open and Distributed Learning, 11(1), 61–80.",
        "Brown, J. S., Collins, A., and Duguid, P. (1989). Situated cognition and the culture of learning. Educational Researcher, 18(1), 32–42.",
        "Gašević, D., Dawson, S., and Siemens, G. (2015). Let’s not forget: Learning analytics are about learning. TechTrends, 59(1), 64–71.",
        "Hevner, A. R., March, S. T., Park, J., and Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75–105.",
        "Knowles, M. S. (1984). Andragogy in action: Applying modern principles of adult learning. San Francisco, CA: Jossey-Bass.",
        "Richardson, J. C., and Swan, K. (2003). Examining social presence in online courses in relation to students’ perceived learning and satisfaction. Journal of Asynchronous Learning Networks, 7(1), 68–88.",
        "Siemens, G., and Baker, R. S. (2012). Learning analytics and educational data mining: towards communication and collaboration. In Proceedings of the 2nd International Conference on Learning Analytics and Knowledge (LAK '12), pp. 252–254.",
        "Sweller, J. (1988). Cognitive load during problem solving: Effects on learning. Cognitive Science, 12(2), 257–285."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.left_indent = Inches(0.3)
        p_ref.paragraph_format.first_line_indent = Inches(-0.3)
        r = p_ref.add_run(ref)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)

    output_path = "data/Context.docx"
    doc.save(output_path)
    print(f"Successfully generated updated paper at: {output_path}")

if __name__ == '__main__':
    build_paper()
